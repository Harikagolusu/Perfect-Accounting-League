"""Workbook generation helpers.

The client workbooks are used as reference templates, not overwritten. They are
copied into the package so the generator can preserve the exact sheet names,
headers, header styling, column widths and Read Me sheet. Only the data rows on
the data sheet are replaced.

The original files supplied by the client stay untouched in the repository
root; generated files are returned as bytes and written elsewhere.
"""

from __future__ import annotations

import io
import re
from copy import copy
from datetime import date
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates"


def template_path(filename: str) -> Path:
    """Return a bundled reference template path, failing loudly if absent."""
    path = TEMPLATE_DIR / filename
    if not path.is_file():
        raise FileNotFoundError(f"Client export template is missing: {path}")
    return path


def load_template(filename: str) -> Workbook:
    return load_workbook(template_path(filename), data_only=False)


def prepare_data_sheet(
    workbook: Workbook,
    sheet_name: str,
    *,
    sample_row: int = 2,
) -> tuple[Worksheet, list[object]]:
    """Clear sample data rows while preserving the header and sheet metadata.

    Returns the sheet and a captured list of cell styles from the first sample
    data row. Those styles are reused so dates and money keep the client's
    number formats.
    """
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Template has no sheet named {sheet_name!r}")

    sheet = workbook[sheet_name]
    styles: list[object] = []
    if sheet.max_row >= sample_row:
        styles = [copy(cell._style) for cell in sheet[sample_row]]
        sheet.delete_rows(sample_row, sheet.max_row - sample_row + 1)

    # A stale autofilter covering only the old sample rows would hide newly
    # written data. The caller may set a fresh range after writing.
    sheet.auto_filter.ref = None
    return sheet, styles


def write_rows(
    sheet: Worksheet,
    rows: Sequence[Sequence[object]],
    *,
    styles: Sequence[object] | None = None,
    number_formats: dict[int, str] | None = None,
    wrap_columns: Iterable[int] = (),
) -> None:
    """Write rows below the header, preserving captured template styles."""
    formats = number_formats or {}
    wrapped = set(wrap_columns)

    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            if styles and column_index - 1 < len(styles):
                cell._style = copy(styles[column_index - 1])
            if column_index in formats:
                cell.number_format = formats[column_index]
            if column_index in wrapped and value is not None:
                cell.alignment = Alignment(
                    horizontal=cell.alignment.horizontal,
                    vertical="top",
                    wrap_text=True,
                )


def set_data_filter(sheet: Worksheet, column_count: int, row_count: int) -> None:
    """Cover the header and every written data row, or just the header."""
    last_column = sheet.cell(row=1, column=column_count).column_letter
    last_row = max(1, row_count + 1)
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"


def workbook_bytes(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def safe_filename_piece(value: str) -> str:
    """Keep generated filenames portable across Windows and browsers."""
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return cleaned.strip("._") or "PAL"


def download_filename(company_name: str, label: str, *, today: date | None = None) -> str:
    stamp = (today or date.today()).strftime("%Y%m%d")
    return f"PAL_{safe_filename_piece(company_name)}_{safe_filename_piece(label)}_{stamp}.xlsx"


def content_disposition(filename: str) -> str:
    # ASCII-only filenames are produced by `safe_filename_piece`, so quoting is
    # sufficient and no RFC 5987 encoding is needed.
    return f'attachment; filename="{filename}"'


__all__ = [
    "TEMPLATE_DIR",
    "content_disposition",
    "download_filename",
    "load_template",
    "prepare_data_sheet",
    "safe_filename_piece",
    "set_data_filter",
    "template_path",
    "workbook_bytes",
    "write_rows",
]
