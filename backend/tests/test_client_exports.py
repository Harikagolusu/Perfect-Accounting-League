"""End-to-end tests for the three client-format exports.

The tests generate workbooks through the API and read them back with openpyxl,
so they verify the file the client actually receives: sheet names, exact column
order, row structure, number formats and company scoping.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.exports.accounting_voucher_export.schemas import ACCOUNTING_COLUMNS
from app.services.exports.bank_export.schemas import BANK_COLUMNS
from app.services.exports.invoice_export.schemas import INVOICE_COLUMNS

COMPANIES_URL = "/api/v1/companies"
PDF_BYTES = b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog >>\nendobj\ntrailer\n%%EOF\n"

INVOICE_BASE: dict[str, Any] = {
    "invoice_number": "INV-001",
    "invoice_date": "2026-03-15",
    "vendor_name": "Deccan Freight Carriers",
    "vendor_gstin": "27AACCD1234F1Z2",
    "vendor_pan": "AACCD1234F",
    "buyer_gstin": "27AAJCS4821K1Z5",
    "place_of_supply": "27-Maharashtra",
    "taxable_amount": "1000.00",
    "cgst_amount": "90.00",
    "sgst_amount": "90.00",
    "cess_amount": "0.00",
    "round_off": "0.00",
    "total_amount": "1180.00",
}

LINE_BASE: dict[str, Any] = {
    "description": "Industrial compressor",
    "supply_type": "GOODS",
    "hsn_sac": "8414",
    "quantity": "2.000",
    "gst_rate": "18.00",
    "taxable_amount": "1000.00",
    "cgst_amount": "90.00",
    "sgst_amount": "90.00",
    "cess_amount": "0.00",
    "total_amount": "1180.00",
}


def _base(company_id: str) -> str:
    return f"{COMPANIES_URL}/{company_id}"


def _company(client: TestClient, name: str = "Test Co", **overrides: Any) -> dict[str, Any]:
    response = client.post(COMPANIES_URL, json={"name": name, **overrides})
    assert response.status_code == 201, response.text
    return response.json()


def _document(client: TestClient, company_id: str, filename: str = "invoice.pdf") -> dict[str, Any]:
    response = client.post(
        f"{_base(company_id)}/documents",
        files={"file": (filename, PDF_BYTES, "application/pdf")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _vendor(client: TestClient, company_id: str, **overrides: Any) -> dict[str, Any]:
    payload: dict[str, Any] = {
        "name": "Deccan Freight Carriers",
        "gstin": "27AACCD1234F1Z2",
        "pan": "AACCD1234F",
        "address": "12 Market Road, Pune",
        "pincode": "411001",
        "msme_registration": "UDYAM-MH-01-0001234",
        "creditor_ledger": "Deccan Freight Carriers",
        "default_expense_ledger": "Freight Inward",
        **overrides,
    }
    response = client.post(f"{_base(company_id)}/vendors", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _invoice(
    client: TestClient,
    company_id: str,
    document_id: str,
    *,
    vendor_id: str | None = None,
    **overrides: Any,
) -> dict[str, Any]:
    payload = {**INVOICE_BASE, "document_id": document_id, **overrides}
    if vendor_id is not None:
        payload["vendor_id"] = vendor_id
    response = client.post(f"{_base(company_id)}/invoices", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _line_item(
    client: TestClient,
    company_id: str,
    invoice_id: str,
    **overrides: Any,
) -> dict[str, Any]:
    response = client.post(
        f"{_base(company_id)}/invoices/{invoice_id}/line-items",
        json={**LINE_BASE, **overrides},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _ready_invoice(
    client: TestClient,
    *,
    company_overrides: dict[str, Any] | None = None,
    vendor_overrides: dict[str, Any] | None = None,
    invoice_overrides: dict[str, Any] | None = None,
    line_overrides: dict[str, Any] | None = None,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any], dict[str, Any]]:
    company = _company(client, "Asterleaf Traders", **(company_overrides or {}))
    document = _document(client, company["id"])
    vendor = _vendor(client, company["id"], **(vendor_overrides or {}))
    invoice = _invoice(
        client,
        company["id"],
        document["id"],
        vendor_id=vendor["id"],
        **(invoice_overrides or {}),
    )
    line = _line_item(
        client, company["id"], invoice["id"], **(line_overrides or {})
    )
    return company, document, invoice, line


def _bank_account(
    client: TestClient,
    company_id: str,
    *,
    ledger_name: str | None = "Canara Bank Current A/c - 6478",
    **overrides: Any,
) -> dict[str, Any]:
    payload = {
        "bank_name": "Canara Bank",
        "account_name": "Current",
        "account_number": "50100234567890",
        "ledger_name": ledger_name,
        **overrides,
    }
    response = client.post(f"{_base(company_id)}/bank-accounts", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _statement(client: TestClient, company_id: str, account_id: str) -> dict[str, Any]:
    response = client.post(
        f"{_base(company_id)}/bank-statements",
        json={
            "bank_account_id": account_id,
            "statement_start_date": "2026-03-01",
            "statement_end_date": "2026-03-31",
            "opening_balance": "10000.00",
            "closing_balance": "12000.00",
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _transaction(
    client: TestClient,
    company_id: str,
    statement_id: str,
    *,
    credit: bool = True,
    counterparty_ledger: str | None = "Acme Customer",
    **overrides: Any,
) -> dict[str, Any]:
    payload = {
        "transaction_date": "2026-03-10",
        "description": "NEFT CR ACME CUSTOMER",
        "debit_amount": "0.00" if credit else "2000.00",
        "credit_amount": "2000.00" if credit else "0.00",
        "counterparty_name": "Acme Customer",
        "counterparty_ledger": counterparty_ledger,
        "counterparty_address": "88 Buyer Street",
        "counterparty_pincode": "560001",
        "counterparty_gstin": "29AAJCS4821K1Z5",
        **overrides,
    }
    response = client.post(
        f"{_base(company_id)}/bank-statements/{statement_id}/transactions",
        json=payload,
    )
    assert response.status_code == 201, response.text
    return response.json()


def _workbook(response) -> Any:
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return load_workbook(io.BytesIO(response.content))


def _data_rows(sheet) -> list[tuple[Any, ...]]:
    rows = [
        tuple(row)
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in row)
    ]
    return rows


def test_invoice_export_matches_client_columns_and_sheet(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(client)

    response = client.get(
        f"{_base(company['id'])}/invoices/export.xlsx",
        params={"invoice_id": invoice["id"]},
    )
    workbook = _workbook(response)

    assert workbook.sheetnames == ["Sheet1"]
    sheet = workbook["Sheet1"]
    assert tuple(cell.value for cell in sheet[1]) == INVOICE_COLUMNS

    rows = _data_rows(sheet)
    assert len(rows) == 1
    row = rows[0]
    assert row[0].date() == date(2026, 3, 15)
    assert row[1] == "INV-001"
    assert row[2] == "Deccan Freight Carriers"
    assert row[4] == "27AACCD1234F1Z2"
    assert row[5] == "AACCD1234F"
    assert row[6] == "27-Maharashtra"
    assert row[8] == "Industrial compressor"
    assert row[9] == "Goods"
    assert row[10] == "8414"
    assert row[11] == 2
    assert row[12] == 18
    assert row[13] == 1000
    assert row[14] is None
    assert row[15] == 90
    assert row[16] == 90
    assert row[17] == 1180


def test_invoice_export_writes_one_row_per_line_item(client: TestClient) -> None:
    company, document, _first_invoice, _line = _ready_invoice(client)
    vendor = _vendor(client, company["id"], name="Second Vendor")
    invoice = _invoice(client, company["id"], document["id"], vendor_id=vendor["id"])
    _line_item(
        client,
        company["id"],
        invoice["id"],
        description="Compressor",
        hsn_sac="8414",
        taxable_amount="600.00",
        total_amount="708.00",
        cgst_amount="54.00",
        sgst_amount="54.00",
    )
    _line_item(
        client,
        company["id"],
        invoice["id"],
        line_number=2,
        description="Installation",
        supply_type="SERVICE",
        hsn_sac="998739",
        quantity="1.000",
        gst_rate="18.00",
        taxable_amount="400.00",
        cgst_amount="36.00",
        sgst_amount="36.00",
        total_amount="472.00",
    )

    response = client.get(
        f"{_base(company['id'])}/invoices/export.xlsx",
        params={"invoice_id": invoice["id"]},
    )
    rows = _data_rows(_workbook(response)["Sheet1"])

    assert len(rows) == 2
    assert [row[8] for row in rows] == ["Compressor", "Installation"]
    assert [row[9] for row in rows] == ["Goods", "Service"]
    assert [row[10] for row in rows] == ["8414", "998739"]
    assert [row[11] for row in rows] == [2, 1]
    assert [row[17] for row in rows] == [708, 472]
    assert all(row[1] == invoice["invoice_number"] for row in rows)


def test_invoice_export_maps_igst_and_leaves_other_taxes_blank(client: TestClient) -> None:
    _company_row, _document_row, _invoice_row, _line = _ready_invoice(
        client,
        invoice_overrides={
            "cgst_amount": None,
            "sgst_amount": None,
            "igst_amount": "180.00",
            "total_amount": "1180.00",
        },
        line_overrides={
            "cgst_amount": None,
            "sgst_amount": None,
            "igst_amount": "180.00",
        },
    )
    company = _company_row
    invoice = _invoice_row
    response = client.get(
        f"{_base(company['id'])}/invoices/export.xlsx",
        params={"invoice_id": invoice["id"]},
    )
    row = _data_rows(_workbook(response)["Sheet1"])[0]
    assert row[14] == 180
    assert row[15] is None
    assert row[16] is None


def test_missing_invoice_number_is_an_exception_not_a_row(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(
        client, invoice_overrides={"invoice_number": None}
    )

    preview = client.get(
        f"{_base(company['id'])}/invoices/export/preview",
        params={"invoice_id": invoice["id"]},
    )
    body = preview.json()
    assert body["rows"] == 0
    assert body["excluded_records"] == 1
    assert any(
        issue["code"] == "INVOICE_FIELD_REQUIRED" and issue["field"] == "Inv Number"
        for issue in body["exceptions"]
    )

    workbook = _workbook(
        client.get(
            f"{_base(company['id'])}/invoices/export.xlsx",
            params={"invoice_id": invoice["id"]},
        )
    )
    assert _data_rows(workbook["Sheet1"]) == []


def test_missing_gstin_is_an_exception(client: TestClient) -> None:
    company = _company(client)
    document = _document(client, company["id"])
    vendor = _vendor(client, company["id"], gstin=None)
    invoice = _invoice(
        client,
        company["id"],
        document["id"],
        vendor_id=vendor["id"],
        vendor_gstin=None,
    )
    _line_item(client, company["id"], invoice["id"])

    body = client.get(
        f"{_base(company['id'])}/invoices/export/preview",
        params={"invoice_id": invoice["id"]},
    ).json()
    assert any(
        issue["code"] == "INVOICE_FIELD_REQUIRED" and issue["field"] == "GST Reg No"
        for issue in body["exceptions"]
    )


def test_missing_place_of_supply_is_an_exception(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(
        client, invoice_overrides={"place_of_supply": None}
    )
    body = client.get(
        f"{_base(company['id'])}/invoices/export/preview",
        params={"invoice_id": invoice["id"]},
    ).json()
    assert any(
        issue["field"] == "Place of Supply" for issue in body["exceptions"]
    )


def test_missing_line_items_is_an_exception(client: TestClient) -> None:
    company = _company(client)
    document = _document(client, company["id"])
    vendor = _vendor(client, company["id"])
    invoice = _invoice(client, company["id"], document["id"], vendor_id=vendor["id"])

    body = client.get(
        f"{_base(company['id'])}/invoices/export/preview",
        params={"invoice_id": invoice["id"]},
    ).json()
    assert body["rows"] == 0
    assert any(
        issue["code"] == "INVOICE_LINE_ITEMS_REQUIRED"
        for issue in body["exceptions"]
    )


def test_missing_gst_rate_is_a_tax_exception(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(
        client, line_overrides={"gst_rate": None}
    )
    body = client.get(
        f"{_base(company['id'])}/invoices/export/preview",
        params={"invoice_id": invoice["id"]},
    ).json()
    assert any(
        issue["code"] == "TAX_DATA_REQUIRED" and issue["field"] == "GST Rate"
        for issue in body["exceptions"]
    )


def test_unmapped_vendor_is_an_exception(client: TestClient) -> None:
    company = _company(client)
    document = _document(client, company["id"])
    invoice = _invoice(client, company["id"], document["id"])
    _line_item(client, company["id"], invoice["id"])

    body = client.get(
        f"{_base(company['id'])}/invoices/export/preview",
        params={"invoice_id": invoice["id"]},
    ).json()
    assert any(
        issue["code"] == "VENDOR_MAPPING_REQUIRED"
        for issue in body["exceptions"]
    )


def test_vendor_master_fields_and_pan_derivation_are_used(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(
        client,
        invoice_overrides={"vendor_pan": None},
    )
    response = client.get(
        f"{_base(company['id'])}/invoices/export.xlsx",
        params={"invoice_id": invoice["id"]},
    )
    row = _data_rows(_workbook(response)["Sheet1"])[0]
    assert row[5] == "AACCD1234F"
    assert row[3] == "UDYAM-MH-01-0001234"
    assert row[7] == "12 Market Road, Pune"


def test_bank_receipt_voucher_matches_client_columns(client: TestClient) -> None:
    company = _company(client)
    account = _bank_account(client, company["id"])
    statement = _statement(client, company["id"], account["id"])
    _transaction(client, company["id"], statement["id"], credit=True)

    response = client.get(
        f"{_base(company['id'])}/bank-transactions/export.xlsx",
        params={"statement_id": statement["id"]},
    )
    workbook = _workbook(response)
    assert workbook.sheetnames == ["Accounting Voucher", "Accounting Voucher (Read Me)"]
    sheet = workbook["Accounting Voucher"]
    assert tuple(cell.value for cell in sheet[1]) == BANK_COLUMNS
    assert workbook["Accounting Voucher (Read Me)"].max_row > 400

    rows = _data_rows(sheet)
    assert len(rows) == 2
    bank_row, party_row = rows
    assert bank_row[0] == "10 Mar 2026"
    assert bank_row[1] == "Receipt"
    assert bank_row[2] == 1
    assert bank_row[5] == "Canara Bank Current A/c - 6478"
    assert bank_row[6] == 2000
    assert bank_row[7] == "Dr"
    assert bank_row[13] == "Accounting"
    assert bank_row[14] == "NEFT CR ACME CUSTOMER"

    assert party_row[5] == "Acme Customer"
    assert party_row[6] == 2000
    assert party_row[7] == "Cr"
    assert party_row[3] == "88 Buyer Street"
    assert party_row[4] == "560001"


def test_bank_payment_voucher_direction_is_reversed(client: TestClient) -> None:
    company = _company(client)
    account = _bank_account(client, company["id"])
    statement = _statement(client, company["id"], account["id"])
    _transaction(client, company["id"], statement["id"], credit=False)

    response = client.get(
        f"{_base(company['id'])}/bank-transactions/export.xlsx",
        params={"statement_id": statement["id"]},
    )
    rows = _data_rows(_workbook(response)["Accounting Voucher"])
    assert rows[0][1] == "Payment"
    assert rows[0][7] == "Cr"
    assert rows[1][7] == "Dr"
    assert rows[0][6] == rows[1][6] == 2000


def test_bank_missing_counterparty_mapping_is_an_exception(client: TestClient) -> None:
    company = _company(client)
    account = _bank_account(client, company["id"])
    statement = _statement(client, company["id"], account["id"])
    _transaction(
        client,
        company["id"],
        statement["id"],
        counterparty_ledger=None,
    )

    body = client.get(
        f"{_base(company['id'])}/bank-transactions/export/preview",
        params={"statement_id": statement["id"]},
    ).json()
    assert body["rows"] == 0
    assert body["excluded_records"] == 1
    assert any(issue["code"] == "BANK_MAPPING_REQUIRED" for issue in body["exceptions"])


def test_bank_missing_account_ledger_is_an_exception(client: TestClient) -> None:
    company = _company(client)
    account = _bank_account(client, company["id"], ledger_name=None)
    statement = _statement(client, company["id"], account["id"])
    _transaction(client, company["id"], statement["id"])

    body = client.get(
        f"{_base(company['id'])}/bank-transactions/export/preview",
        params={"statement_id": statement["id"]},
    ).json()
    assert any(
        issue["code"] == "LEDGER_MAPPING_REQUIRED"
        for issue in body["exceptions"]
    )


def test_bank_voucher_numbers_restart_per_voucher_type(client: TestClient) -> None:
    company = _company(client)
    account = _bank_account(client, company["id"])
    statement = _statement(client, company["id"], account["id"])
    _transaction(client, company["id"], statement["id"], credit=True)
    _transaction(client, company["id"], statement["id"], credit=True)
    _transaction(client, company["id"], statement["id"], credit=False)

    response = client.get(
        f"{_base(company['id'])}/bank-transactions/export.xlsx",
        params={"statement_id": statement["id"]},
    )
    rows = _data_rows(_workbook(response)["Accounting Voucher"])
    receipts = [row for row in rows if row[1] == "Receipt"]
    payments = [row for row in rows if row[1] == "Payment"]
    assert [row[2] for row in receipts] == [1, 1, 2, 2]
    assert [row[2] for row in payments] == [1, 1]


def test_bank_partial_export_keeps_valid_rows_and_reports_exception(
    client: TestClient,
) -> None:
    company = _company(client)
    account = _bank_account(client, company["id"])
    statement = _statement(client, company["id"], account["id"])
    _transaction(client, company["id"], statement["id"], credit=True)
    _transaction(
        client,
        company["id"],
        statement["id"],
        credit=False,
        counterparty_ledger=None,
    )

    preview = client.get(
        f"{_base(company['id'])}/bank-transactions/export/preview",
        params={"statement_id": statement["id"]},
    ).json()
    assert preview["rows"] == 2
    assert preview["excluded_records"] == 1
    assert len(preview["exceptions"]) == 1

    workbook = _workbook(
        client.get(
            f"{_base(company['id'])}/bank-transactions/export.xlsx",
            params={"statement_id": statement["id"]},
        )
    )
    assert len(_data_rows(workbook["Accounting Voucher"])) == 2


def test_accounting_voucher_from_purchase_invoice_is_balanced(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(
        client,
        company_overrides={
            "cgst_ledger": "CGST",
            "sgst_ledger": "SGST",
            "igst_ledger": "IGST",
            "round_off_ledger": "Round Off",
        },
        invoice_overrides={
            "round_off": "0.50",
            "total_amount": "1180.50",
        },
        line_overrides={"total_amount": "1180.50"},
    )

    response = client.get(
        f"{_base(company['id'])}/accounting-vouchers/export.xlsx",
        params={"include_bank": "false"},
    )
    workbook = _workbook(response)
    assert workbook.sheetnames == ["Accounting Voucher", "Accounting Voucher (Read Me)"]
    sheet = workbook["Accounting Voucher"]
    assert tuple(cell.value for cell in sheet[1]) == ACCOUNTING_COLUMNS
    rows = _data_rows(sheet)

    assert len(rows) == 5
    assert {row[1] for row in rows} == {"Purchase"}
    assert {row[16] for row in rows} == {"INV-001"}
    debit = sum(row[6] for row in rows if row[7] == "Dr")
    credit = sum(row[6] for row in rows if row[7] == "Cr")
    assert debit == credit == 1180.5
    assert [row[5] for row in rows] == [
        "Deccan Freight Carriers",
        "Freight Inward",
        "CGST",
        "SGST",
        "Round Off",
    ]


def test_accounting_voucher_unbalanced_invoice_is_an_exception(
    client: TestClient,
) -> None:
    company, _document_row, invoice, _line = _ready_invoice(
        client,
        company_overrides={
            "cgst_ledger": "CGST",
            "sgst_ledger": "SGST",
            "igst_ledger": "IGST",
        },
        invoice_overrides={"total_amount": "1200.00"},
        line_overrides={"total_amount": "1200.00"},
    )
    body = client.get(
        f"{_base(company['id'])}/accounting-vouchers/export/preview",
        params={"include_bank": "false"},
    ).json()
    assert body["rows"] == 0
    assert any(issue["code"] == "VOUCHER_NOT_BALANCED" for issue in body["exceptions"])


def test_accounting_voucher_missing_invoice_total_is_an_exception(
    client: TestClient,
) -> None:
    company, _document_row, invoice, _line = _ready_invoice(
        client,
        company_overrides={
            "cgst_ledger": "CGST",
            "sgst_ledger": "SGST",
            "igst_ledger": "IGST",
        },
        invoice_overrides={"total_amount": None},
    )

    body = client.get(
        f"{_base(company['id'])}/accounting-vouchers/export/preview",
        params={"include_bank": "false"},
    ).json()

    assert body["rows"] == 0
    assert any(
        issue["code"] == "INVOICE_FIELD_REQUIRED"
        and issue["field"] == "Ledger Amount"
        and "total" in issue["message"].lower()
        for issue in body["exceptions"]
    )
    assert not any(
        issue["code"] == "VOUCHER_NOT_BALANCED" for issue in body["exceptions"]
    )


def test_accounting_voucher_missing_gst_ledger_is_an_exception(
    client: TestClient,
) -> None:
    company, _document_row, invoice, _line = _ready_invoice(client)
    body = client.get(
        f"{_base(company['id'])}/accounting-vouchers/export/preview",
        params={"include_bank": "false"},
    ).json()
    assert any(
        issue["code"] == "LEDGER_MAPPING_REQUIRED"
        and "CGST ledger" in issue["message"]
        for issue in body["exceptions"]
    )


def test_accounting_voucher_can_include_bank_transactions(client: TestClient) -> None:
    company = _company(
        client,
        cgst_ledger="CGST",
        sgst_ledger="SGST",
        igst_ledger="IGST",
    )
    account = _bank_account(client, company["id"])
    statement = _statement(client, company["id"], account["id"])
    _transaction(client, company["id"], statement["id"], credit=True)

    body = client.get(
        f"{_base(company['id'])}/accounting-vouchers/export/preview",
        params={"include_invoices": "false", "include_bank": "true"},
    ).json()
    assert body["rows"] == 2
    assert body["vouchers"] == 1

    workbook = _workbook(
        client.get(
            f"{_base(company['id'])}/accounting-vouchers/export.xlsx",
            params={"include_invoices": "false", "include_bank": "true"},
        )
    )
    rows = _data_rows(workbook["Accounting Voucher"])
    assert rows[0][1] == "Receipt"
    assert rows[0][13] == "Accounting"
    assert rows[0][16] is None


def test_invoice_export_is_company_scoped(client: TestClient) -> None:
    company_a, _document_a, _invoice_a, _line_a = _ready_invoice(
        client, invoice_overrides={"invoice_number": "ONLY-A"}
    )
    company_b = _company(client, "Other Co")
    document_b = _document(client, company_b["id"], "b.pdf")
    vendor_b = _vendor(client, company_b["id"], name="Other Vendor")
    invoice_b = _invoice(
        client,
        company_b["id"],
        document_b["id"],
        vendor_id=vendor_b["id"],
        invoice_number="ONLY-B",
    )
    _line_item(client, company_b["id"], invoice_b["id"])

    response = client.get(f"{_base(company_a['id'])}/invoices/export.xlsx")
    workbook = _workbook(response)
    rows = _data_rows(workbook["Sheet1"])
    assert [row[1] for row in rows] == ["ONLY-A"]
    assert "ONLY-B" not in response.content.decode("latin-1")


def test_bank_export_is_company_scoped(client: TestClient) -> None:
    company_a = _company(client, "Company A")
    account_a = _bank_account(client, company_a["id"])
    statement_a = _statement(client, company_a["id"], account_a["id"])
    _transaction(
        client,
        company_a["id"],
        statement_a["id"],
        counterparty_ledger="A Party",
    )

    company_b = _company(client, "Company B")
    account_b = _bank_account(client, company_b["id"], ledger_name="B Bank")
    statement_b = _statement(client, company_b["id"], account_b["id"])
    _transaction(
        client,
        company_b["id"],
        statement_b["id"],
        counterparty_ledger="B Party",
    )

    response = client.get(f"{_base(company_a['id'])}/bank-transactions/export.xlsx")
    rows = _data_rows(_workbook(response)["Accounting Voucher"])
    ledger_names = {row[5] for row in rows}
    assert "A Party" in ledger_names
    assert "B Party" not in ledger_names


def test_exports_are_idempotent_for_the_same_records(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(client)
    url = f"{_base(company['id'])}/invoices/export/preview"

    first = client.get(url, params={"invoice_id": invoice["id"]}).json()
    second = client.get(url, params={"invoice_id": invoice["id"]}).json()
    assert first == second

    workbook_a = _workbook(
        client.get(
            f"{_base(company['id'])}/invoices/export.xlsx",
            params={"invoice_id": invoice["id"]},
        )
    )
    workbook_b = _workbook(
        client.get(
            f"{_base(company['id'])}/invoices/export.xlsx",
            params={"invoice_id": invoice["id"]},
        )
    )
    assert _data_rows(workbook_a["Sheet1"]) == _data_rows(workbook_b["Sheet1"])


def test_empty_export_still_has_the_client_header(client: TestClient) -> None:
    company = _company(client)
    response = client.get(f"{_base(company['id'])}/invoices/export.xlsx")
    workbook = _workbook(response)
    assert _data_rows(workbook["Sheet1"]) == []
    assert tuple(cell.value for cell in workbook["Sheet1"][1]) == INVOICE_COLUMNS


def test_invoice_date_and_money_number_formats_are_set(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(client)
    workbook = _workbook(
        client.get(
            f"{_base(company['id'])}/invoices/export.xlsx",
            params={"invoice_id": invoice["id"]},
        )
    )
    sheet = workbook["Sheet1"]
    assert sheet.cell(row=2, column=1).number_format == "dd/mm/yyyy"
    assert sheet.cell(row=2, column=14).number_format == "#,##0.00"
    assert sheet.cell(row=2, column=18).number_format == "#,##0.00"


def test_line_item_api_auto_numbers_and_is_company_scoped(client: TestClient) -> None:
    company, document, _first_invoice, _line = _ready_invoice(client)
    vendor = _vendor(client, company["id"], name="Second Vendor")
    invoice = _invoice(client, company["id"], document["id"], vendor_id=vendor["id"])
    first = _line_item(client, company["id"], invoice["id"])
    second = _line_item(client, company["id"], invoice["id"], description="Second line")
    assert first["line_number"] == 1
    assert second["line_number"] == 2

    other = _company(client, "Other Co")
    response = client.get(
        f"{_base(other['id'])}/invoices/{invoice['id']}/line-items"
    )
    assert response.status_code == 404


def test_line_item_supply_type_is_normalised(client: TestClient) -> None:
    company, document, _first_invoice, _line = _ready_invoice(client)
    vendor = _vendor(client, company["id"], name="Normalised Vendor")
    invoice = _invoice(client, company["id"], document["id"], vendor_id=vendor["id"])

    line = _line_item(client, company["id"], invoice["id"], supply_type="goods")

    assert line["supply_type"] == "GOODS"


def test_bank_mapping_endpoint_changes_only_mapping(client: TestClient) -> None:
    company = _company(client)
    account = _bank_account(client, company["id"])
    statement = _statement(client, company["id"], account["id"])
    transaction = _transaction(
        client,
        company["id"],
        statement["id"],
        counterparty_ledger=None,
    )

    response = client.patch(
        f"{_base(company['id'])}/bank-statements/{statement['id']}"
        f"/transactions/{transaction['id']}/mapping",
        json={
            "counterparty_name": "Fixed Party",
            "counterparty_ledger": "Fixed Party Ledger",
            "counterparty_address": "1 Fixed Road",
            "counterparty_pincode": "400001",
            "counterparty_gstin": "27AAJCS4821K1Z5",
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["counterparty_ledger"] == "Fixed Party Ledger"
    assert body["credit_amount"] == "2000.00"
    assert body["debit_amount"] == "0.00"
    assert body["description"] == "NEFT CR ACME CUSTOMER"


def test_accounting_voucher_date_is_written_as_a_date_cell(client: TestClient) -> None:
    company, _document_row, invoice, _line = _ready_invoice(
        client,
        company_overrides={
            "cgst_ledger": "CGST",
            "sgst_ledger": "SGST",
            "igst_ledger": "IGST",
        },
    )
    workbook = _workbook(
        client.get(
            f"{_base(company['id'])}/accounting-vouchers/export.xlsx",
            params={"include_bank": "false"},
        )
    )
    cell = workbook["Accounting Voucher"].cell(row=2, column=1)
    assert isinstance(cell.value, datetime)
    assert cell.value.date() == date(2026, 3, 15)
    assert cell.number_format == "dd/mm/yyyy"


def test_export_unknown_company_returns_404(client: TestClient) -> None:
    missing = "00000000-0000-0000-0000-000000000000"
    assert client.get(f"{_base(missing)}/invoices/export/preview").status_code == 404
    assert client.get(f"{_base(missing)}/bank-transactions/export/preview").status_code == 404
    assert (
        client.get(f"{_base(missing)}/accounting-vouchers/export/preview").status_code
        == 404
    )


if __name__ == "__main__":  # pragma: no cover
    pytest.main([__file__])
