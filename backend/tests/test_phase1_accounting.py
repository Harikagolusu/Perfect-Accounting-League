"""Phase 1 accounting workflow tests.

Covers the ledger master, explainable mappings, draft vouchers, approval,
audit, invoice/bank creation, export lifecycle and cross-company isolation.
"""

from __future__ import annotations

import io
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.services.exports.accounting_voucher_export.schemas import ACCOUNTING_COLUMNS

COMPANIES_URL = "/api/v1/companies"
PDF_BYTES = b"%PDF-1.7\n1 0 obj\n<< /Type /Catalog >>\nendobj\ntrailer\n%%EOF\n"


def _base(company_id: str) -> str:
    return f"{COMPANIES_URL}/{company_id}"


def _company(client: TestClient, name: str = "Phase One Co", **overrides: Any) -> dict[str, Any]:
    payload = {
        "name": name,
        "cgst_ledger": "CGST Input",
        "sgst_ledger": "SGST Input",
        "igst_ledger": "IGST Input",
        "cess_ledger": "Cess Input",
        "round_off_ledger": "Round Off",
        "sales_ledger": "Sales",
        **overrides,
    }
    response = client.post(COMPANIES_URL, json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _ledger(client: TestClient, company_id: str, name: str, **overrides: Any) -> dict[str, Any]:
    response = client.post(
        f"{_base(company_id)}/ledgers",
        json={"name": name, "ledger_type": "Test", **overrides},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _document(client: TestClient, company_id: str, name: str = "invoice.pdf") -> dict[str, Any]:
    response = client.post(
        f"{_base(company_id)}/documents",
        files={"file": (name, PDF_BYTES, "application/pdf")},
    )
    assert response.status_code == 201, response.text
    return response.json()


def _vendor(
    client: TestClient,
    company_id: str,
    *,
    name: str = "Deccan Freight Carriers",
    expense_ledger: str | None = "Freight Inward",
    creditor_ledger: str | None = "Deccan Freight Carriers",
) -> dict[str, Any]:
    response = client.post(
        f"{_base(company_id)}/vendors",
        json={
            "name": name,
            "gstin": "27AACCD1234F1Z2",
            "pan": "AACCD1234F",
            "address": "12 Market Road, Pune",
            "pincode": "411001",
            "default_expense_ledger": expense_ledger,
            "creditor_ledger": creditor_ledger,
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def _invoice(
    client: TestClient,
    company_id: str,
    document_id: str,
    vendor_id: str,
    **overrides: Any,
) -> dict[str, Any]:
    payload = {
        "document_id": document_id,
        "vendor_id": vendor_id,
        "invoice_number": "INV-001",
        "invoice_date": "2026-03-15",
        "vendor_name": "Deccan Freight Carriers",
        "vendor_gstin": "27AACCD1234F1Z2",
        "vendor_pan": "AACCD1234F",
        "buyer_gstin": "27AAJCS4821K1Z5",
        "place_of_supply": "27-Maharashtra",
        "taxable_amount": "1000.00",
        "cgst_amount": "90.00",
        "sgst_amount": "90.00",
        "cess_amount": "0.00",
        "round_off": "0.00",
        "total_amount": "1180.00",
        "direction": "PURCHASE",
        **overrides,
    }
    response = client.post(f"{_base(company_id)}/invoices", json=payload)
    assert response.status_code == 201, response.text
    return response.json()


def _company_with_ledgers(
    client: TestClient,
    *,
    name: str = "Phase One Co",
    expense: bool = True,
    party: bool = True,
    bank: bool = True,
    taxes: bool = True,
    sales: bool = True,
) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    company = _company(client, name)
    ledgers: dict[str, dict[str, Any]] = {}
    if expense:
        ledgers["expense"] = _ledger(client, company["id"], "Freight Inward")
    if party:
        ledgers["party"] = _ledger(client, company["id"], "Deccan Freight Carriers")
    if bank:
        ledgers["bank"] = _ledger(client, company["id"], "HDFC Current Account")
    if taxes:
        for tax in ("CGST Input", "SGST Input", "IGST Input", "Cess Input"):
            ledgers[tax] = _ledger(client, company["id"], tax)
    if sales:
        ledgers["sales"] = _ledger(client, company["id"], "Sales")
    ledgers["round_off"] = _ledger(client, company["id"], "Round Off")
    return company, ledgers


def _manual_voucher(
    client: TestClient,
    company_id: str,
    debit_id: str,
    credit_id: str,
    *,
    amount: str = "100.00",
    **overrides: Any,
) -> Any:
    payload = {
        "voucher_date": "2026-04-01",
        "voucher_type": "JOURNAL",
        "narration": "Manual adjustment",
        "lines": [
            {"ledger_id": debit_id, "side": "Dr", "amount": amount},
            {"ledger_id": credit_id, "side": "Cr", "amount": amount},
        ],
        **overrides,
    }
    return client.post(f"{_base(company_id)}/vouchers", json=payload)


def _bank_ready(
    client: TestClient,
    company_id: str,
    *,
    ledger_name: str = "HDFC Current Account",
    counterparty_ledger: str | None = "Deccan Freight Carriers",
    credit: str = "500.00",
    debit: str = "0.00",
    **transaction_overrides: Any,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    account = client.post(
        f"{_base(company_id)}/bank-accounts",
        json={
            "bank_name": "HDFC Bank",
            "account_name": "Current Account",
            "account_number": "123456789012",
            "ledger_name": ledger_name,
        },
    )
    assert account.status_code == 201, account.text
    statement = client.post(
        f"{_base(company_id)}/bank-statements",
        json={
            "bank_account_id": account.json()["id"],
            "statement_start_date": "2026-03-01",
            "statement_end_date": "2026-03-31",
            "opening_balance": "1000.00",
            "closing_balance": "1500.00",
        },
    )
    assert statement.status_code == 201, statement.text
    transaction = client.post(
        f"{_base(company_id)}/bank-statements/{statement.json()['id']}/transactions",
        json={
            "transaction_date": "2026-03-20",
            "description": "NEFT from Deccan Freight Carriers",
            "reference_number": "UTR-1",
            "debit_amount": debit,
            "credit_amount": credit,
            "counterparty_name": "Deccan Freight Carriers",
            "counterparty_ledger": counterparty_ledger,
            **transaction_overrides,
        },
    )
    assert transaction.status_code == 201, transaction.text
    return account.json(), statement.json(), transaction.json()


# --- Ledger master ---------------------------------------------------------


def test_ledger_master_crud_is_company_scoped(client: TestClient) -> None:
    a = _company(client, "Company A")
    b = _company(client, "Company B")
    ledger = _ledger(client, a["id"], "Freight Inward")

    listed = client.get(f"{_base(a['id'])}/ledgers")
    assert listed.status_code == 200
    assert [item["id"] for item in listed.json()] == [ledger["id"]]

    other = client.get(f"{_base(b['id'])}/ledgers")
    assert other.status_code == 200
    assert other.json() == []

    cross = client.get(f"{_base(b['id'])}/ledgers/{ledger['id']}")
    assert cross.status_code == 404

    patched = client.patch(
        f"{_base(a['id'])}/ledgers/{ledger['id']}",
        json={"parent_group": "Indirect Expenses"},
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["parent_group"] == "Indirect Expenses"

    deleted = client.delete(f"{_base(a['id'])}/ledgers/{ledger['id']}")
    assert deleted.status_code == 200, deleted.text
    assert deleted.json()["is_active"] is False
    assert client.get(f"{_base(a['id'])}/ledgers").json() == []


def test_ledger_duplicate_name_is_rejected(client: TestClient) -> None:
    company = _company(client)
    _ledger(client, company["id"], "Freight Inward")
    duplicate = client.post(
        f"{_base(company['id'])}/ledgers",
        json={"name": "freight inward"},
    )
    assert duplicate.status_code == 409


def test_ledger_same_name_is_allowed_in_another_company(client: TestClient) -> None:
    a = _company(client, "A")
    b = _company(client, "B")
    _ledger(client, a["id"], "Freight Inward")
    _ledger(client, b["id"], "Freight Inward")


# --- Manual vouchers, approval and audit ----------------------------------


def test_balanced_manual_voucher_can_be_submitted_and_approved(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    created = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    )
    assert created.status_code == 201, created.text
    voucher = created.json()
    assert voucher["status"] == "DRAFT"
    assert voucher["is_balanced"] is True
    assert voucher["total_debit"] == "100.00"
    assert voucher["total_credit"] == "100.00"
    assert voucher["approved_by"] is None

    submitted = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/submit",
        json={"actor": "Asha"},
    )
    assert submitted.status_code == 200, submitted.text
    assert submitted.json()["status"] == "PENDING_APPROVAL"

    approved = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/approve",
        json={"actor": "Asha", "comment": "Reviewed against the invoice."},
    )
    assert approved.status_code == 200, approved.text
    body = approved.json()
    assert body["status"] == "APPROVED"
    assert body["approved_by"] == "Asha"
    assert body["approved_at"] is not None
    assert body["posted_at"] is None

    audit = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    )
    assert audit.status_code == 200, audit.text
    events = audit.json()
    types = [event["event_type"] for event in events]
    assert "VOUCHER_APPROVED" in types
    submission = next(
        event for event in events if event["event_type"] == "VOUCHER_SUBMITTED_FOR_APPROVAL"
    )
    assert submission["actor"] == "Asha"
    approval = next(event for event in events if event["event_type"] == "VOUCHER_APPROVED")
    assert approval["actor"] == "Asha"
    assert approval["previous_status"] == "PENDING_APPROVAL"
    assert approval["new_status"] == "APPROVED"
    assert approval["reason"] == "Reviewed against the invoice."


def test_unbalanced_voucher_routes_to_needs_review_and_cannot_be_approved(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    created = _manual_voucher(
        client,
        company["id"],
        ledgers["expense"]["id"],
        ledgers["party"]["id"],
        amount="100.00",
    )
    voucher = created.json()
    # Replace the lines with an unbalanced pair.
    patched = client.patch(
        f"{_base(company['id'])}/vouchers/{voucher['id']}",
        json={
            "lines": [
                {"ledger_id": ledgers["expense"]["id"], "side": "Dr", "amount": "100.00"},
                {"ledger_id": ledgers["party"]["id"], "side": "Cr", "amount": "90.00"},
            ]
        },
    )
    assert patched.status_code == 200, patched.text
    body = patched.json()
    assert body["status"] == "NEEDS_REVIEW"
    assert body["is_balanced"] is False
    assert any(issue["code"] == "VOUCHER_NOT_BALANCED" for issue in body["validation_issues"])

    submitted = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/submit"
    )
    assert submitted.status_code == 200
    assert submitted.json()["status"] == "NEEDS_REVIEW"

    approved = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/approve",
        json={},
    )
    assert approved.status_code == 409
    detail = approved.json()["detail"]
    assert detail["code"] == "INVALID_STATUS_TRANSITION"


def test_approval_revalidates_and_blocks_when_a_ledger_becomes_inactive(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    submitted = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/submit"
    )
    assert submitted.status_code == 200, submitted.text
    assert submitted.json()["status"] == "PENDING_APPROVAL"

    deactivated = client.delete(
        f"{_base(company['id'])}/ledgers/{ledgers['expense']['id']}"
    )
    assert deactivated.status_code == 200, deactivated.text

    approved = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/approve", json={}
    )
    assert approved.status_code == 409, approved.text
    assert approved.json()["detail"]["code"] == "LEDGER_INACTIVE"
    current = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}"
    ).json()
    assert current["status"] == "NEEDS_REVIEW"
    audit = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    assert any(
        event["event_type"] == "VOUCHER_APPROVAL_BLOCKED"
        and event["previous_status"] == "PENDING_APPROVAL"
        and event["new_status"] == "NEEDS_REVIEW"
        for event in audit
    )


def test_draft_cannot_jump_to_posted_or_export(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()

    exported = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/export"
    )
    assert exported.status_code == 409
    assert exported.json()["detail"]["code"] == "APPROVAL_REQUIRED"

    verified = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/verify-posting",
        json={"verified": True, "external_reference": "TALLY-1"},
    )
    assert verified.status_code == 409
    assert verified.json()["detail"]["code"] == "INVALID_STATUS_TRANSITION"


def test_approval_requires_pending_status(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    response = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/approve",
        json={},
    )
    assert response.status_code == 409
    assert response.json()["detail"]["code"] == "INVALID_STATUS_TRANSITION"


def test_rejected_voucher_can_be_reopened_and_corrected(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    assert client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/submit"
    ).status_code == 200
    rejected = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/reject",
        json={"actor": "Ravi", "reason": "Wrong expense ledger."},
    )
    assert rejected.status_code == 200, rejected.text
    assert rejected.json()["status"] == "REJECTED"
    assert rejected.json()["rejection_reason"] == "Wrong expense ledger."

    rejected_events = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    rejection = next(
        event for event in rejected_events if event["event_type"] == "VOUCHER_REJECTED"
    )
    assert rejection["actor"] == "Ravi"
    assert rejection["reason"] == "Wrong expense ledger."
    assert rejection["previous_status"] == "PENDING_APPROVAL"
    assert rejection["new_status"] == "REJECTED"

    reopened = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/reopen"
    )
    assert reopened.status_code == 200, reopened.text
    assert reopened.json()["status"] == "DRAFT"
    reopened_events = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    assert any(
        event["event_type"] == "VOUCHER_REOPENED"
        and event["previous_status"] == "REJECTED"
        and event["new_status"] == "DRAFT"
        for event in reopened_events
    )


def test_rejection_requires_a_non_blank_reason(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    assert client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/submit"
    ).status_code == 200

    for payload in ({}, {"reason": "   "}, {"comment": ""}):
        rejected = client.post(
            f"{_base(company['id'])}/vouchers/{voucher['id']}/reject",
            json=payload,
        )
        assert rejected.status_code == 409, rejected.text
        assert rejected.json()["detail"]["code"] == "REJECTION_REASON_REQUIRED"

    unchanged = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}"
    ).json()
    assert unchanged["status"] == "PENDING_APPROVAL"
    assert unchanged["rejection_reason"] is None

    accepted = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/reject",
        json={"reason": "  Wrong expense ledger.  "},
    )
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["rejection_reason"] == "Wrong expense ledger."


def test_needs_review_and_posting_failed_vouchers_can_be_reopened(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)

    needs_review = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    patched = client.patch(
        f"{_base(company['id'])}/vouchers/{needs_review['id']}",
        json={
            "lines": [
                {"ledger_id": ledgers["expense"]["id"], "side": "Dr", "amount": "100.00"},
                {"ledger_id": ledgers["party"]["id"], "side": "Cr", "amount": "90.00"},
            ]
        },
    )
    assert patched.status_code == 200, patched.text
    assert patched.json()["status"] == "NEEDS_REVIEW"
    reopened_review = client.post(
        f"{_base(company['id'])}/vouchers/{needs_review['id']}/reopen"
    )
    assert reopened_review.status_code == 200, reopened_review.text
    assert reopened_review.json()["status"] == "DRAFT"

    failed = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    assert client.post(
        f"{_base(company['id'])}/vouchers/{failed['id']}/submit"
    ).status_code == 200
    assert client.post(
        f"{_base(company['id'])}/vouchers/{failed['id']}/approve", json={}
    ).status_code == 200
    assert client.post(
        f"{_base(company['id'])}/vouchers/{failed['id']}/export"
    ).status_code == 200
    posting_failed = client.post(
        f"{_base(company['id'])}/vouchers/{failed['id']}/verify-posting",
        json={"verified": False, "reason": "Tally rejected the voucher."},
    )
    assert posting_failed.status_code == 200, posting_failed.text
    assert posting_failed.json()["status"] == "POSTING_FAILED"
    reopened_failed = client.post(
        f"{_base(company['id'])}/vouchers/{failed['id']}/reopen"
    )
    assert reopened_failed.status_code == 200, reopened_failed.text
    assert reopened_failed.json()["status"] == "DRAFT"


def test_voucher_edit_and_ledger_change_are_audited(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    other_expense = _ledger(client, company["id"], "Other Expense")
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    patched = client.patch(
        f"{_base(company['id'])}/vouchers/{voucher['id']}",
        json={
            "narration": "Corrected narration",
            "lines": [
                {"ledger_id": other_expense["id"], "side": "Dr", "amount": "100.00"},
                {"ledger_id": ledgers["party"]["id"], "side": "Cr", "amount": "100.00"},
            ],
        },
    )
    assert patched.status_code == 200, patched.text
    events = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    types = {event["event_type"] for event in events}
    assert "VOUCHER_UPDATED" in types
    assert "VOUCHER_LINES_REPLACED" in types
    assert "LEDGER_CHANGED" in types


# --- Invoice -> draft voucher ----------------------------------------------


def test_invoice_creates_balanced_purchase_draft_with_configured_ledgers(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    vendor = _vendor(client, company["id"])
    document = _document(client, company["id"])
    invoice = _invoice(client, company["id"], document["id"], vendor["id"])

    response = client.post(
        f"{_base(company['id'])}/vouchers/from-invoice/{invoice['id']}"
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert voucher["source_type"] == "INVOICE"
    assert voucher["source_id"] == invoice["id"]
    assert voucher["voucher_type"] == "PURCHASE"
    assert voucher["is_balanced"] is True
    assert voucher["status"] == "DRAFT"
    names = {line["ledger_name"] for line in voucher["lines"]}
    assert names == {"Freight Inward", "CGST Input", "SGST Input", "Deccan Freight Carriers"}

    audit = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    created = next(
        event
        for event in audit
        if event["event_type"] == "VOUCHER_CREATED_FROM_INVOICE"
    )
    assert created["event_metadata"]["source_id"] == invoice["id"]
    assert created["event_metadata"]["source_document_id"] == document["id"]
    assert created["event_metadata"]["source_values"]["invoice_number"] == "INV-001"

    again = client.post(
        f"{_base(company['id'])}/vouchers/from-invoice/{invoice['id']}"
    )
    assert again.status_code == 200
    assert again.json()["id"] == voucher["id"]


def test_invoice_missing_ledger_routes_to_review_without_inventing_one(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client, expense=False)
    vendor = _vendor(client, company["id"])
    document = _document(client, company["id"])
    invoice = _invoice(client, company["id"], document["id"], vendor["id"])

    response = client.post(
        f"{_base(company['id'])}/vouchers/from-invoice/{invoice['id']}"
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert voucher["status"] == "NEEDS_REVIEW"
    assert any(
        issue["code"] in {"MISSING_LEDGER", "NO_HISTORICAL_MAPPING"}
        for issue in voucher["validation_issues"]
    )
    expense_line = next(line for line in voucher["lines"] if line["side"] == "Dr" and line["line_type"] == "LEDGER")
    assert expense_line["ledger_id"] is None
    assert expense_line["ledger_name"] == "Freight Inward"

    approved = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/approve", json={}
    )
    assert approved.status_code == 409
    assert approved.json()["detail"]["code"] == "INVALID_STATUS_TRANSITION"


def test_ambiguous_mapping_routes_to_review_without_choosing_a_ledger(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    other_expense = _ledger(client, company["id"], "Other Freight Expense")
    vendor = _vendor(client, company["id"], expense_ledger=None)
    document = _document(client, company["id"])
    invoice = _invoice(client, company["id"], document["id"], vendor["id"])

    for ledger in (ledgers["expense"], other_expense):
        mapping = client.post(
            f"{_base(company['id'])}/accounting-mappings",
            json={
                "source_type": "INVOICE",
                "match_type": "VENDOR",
                "vendor_id": vendor["id"],
                "voucher_type": "PURCHASE",
                "ledger_id": ledger["id"],
                "reason": "Two historical choices need a human decision.",
            },
        )
        assert mapping.status_code == 201, mapping.text

    suggestion = client.get(
        f"{_base(company['id'])}/invoices/{invoice['id']}/ledger-suggestion"
    )
    assert suggestion.status_code == 200, suggestion.text
    body = suggestion.json()
    assert any(
        issue["code"] == "AMBIGUOUS_ACCOUNTING_MAPPING"
        for issue in body["exceptions"]
    )
    assert not any(item["purpose"] == "EXPENSE" for item in body["suggestions"])

    response = client.post(
        f"{_base(company['id'])}/vouchers/from-invoice/{invoice['id']}"
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert voucher["status"] == "NEEDS_REVIEW"
    assert any(
        issue["code"] == "AMBIGUOUS_ACCOUNTING_MAPPING"
        for issue in voucher["validation_issues"]
    )
    expense_line = next(
        line
        for line in voucher["lines"]
        if line["side"] == "Dr" and line["line_type"] == "LEDGER"
    )
    assert expense_line["ledger_id"] is None
    assert expense_line["ledger_name"] == ""


def test_invoice_missing_tax_configuration_is_not_guessed(
    client: TestClient,
) -> None:
    company = _company(client, "No CGST Configuration Co", cgst_ledger=None)
    _ledger(client, company["id"], "Freight Inward")
    _ledger(client, company["id"], "Deccan Freight Carriers")
    for tax in ("SGST Input", "IGST Input", "Cess Input"):
        _ledger(client, company["id"], tax)
    vendor = _vendor(client, company["id"])
    document = _document(client, company["id"])
    invoice = _invoice(client, company["id"], document["id"], vendor["id"])

    response = client.post(
        f"{_base(company['id'])}/vouchers/from-invoice/{invoice['id']}"
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert voucher["status"] == "NEEDS_REVIEW"
    assert any(
        issue["code"] == "MISSING_TAX_DATA"
        and issue["field"] == "CGST Ledger"
        for issue in voucher["validation_issues"]
    )
    cgst_line = next(
        line
        for line in voucher["lines"]
        if line["line_type"] == "TAX" and line["ledger_name"] == ""
    )
    assert cgst_line["ledger_id"] is None


def test_invoice_sales_direction_creates_sales_voucher(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    vendor = _vendor(client, company["id"], name="Acme Customer", expense_ledger=None)
    document = _document(client, company["id"])
    invoice = _invoice(
        client,
        company["id"],
        document["id"],
        vendor["id"],
        direction="SALES",
        invoice_number="S-1",
    )
    response = client.post(
        f"{_base(company['id'])}/vouchers/from-invoice/{invoice['id']}"
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert voucher["voucher_type"] == "SALES"
    assert voucher["is_balanced"] is True
    assert any(line["ledger_name"] == "Sales" and line["side"] == "Cr" for line in voucher["lines"])


def test_ledger_suggestion_uses_approved_mapping_after_approval(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    vendor = _vendor(client, company["id"])
    document = _document(client, company["id"])
    first = _invoice(client, company["id"], document["id"], vendor["id"], invoice_number="INV-1")
    voucher = client.post(
        f"{_base(company['id'])}/vouchers/from-invoice/{first['id']}"
    ).json()
    client.post(f"{_base(company['id'])}/vouchers/{voucher['id']}/submit")
    client.post(f"{_base(company['id'])}/vouchers/{voucher['id']}/approve", json={})

    second = _invoice(client, company["id"], document["id"], vendor["id"], invoice_number="INV-2")
    suggestion = client.get(
        f"{_base(company['id'])}/invoices/{second['id']}/ledger-suggestion"
    )
    assert suggestion.status_code == 200, suggestion.text
    body = suggestion.json()
    expense = next(item for item in body["suggestions"] if item["purpose"] == "EXPENSE")
    assert expense["ledger_name"] == "Freight Inward"
    assert expense["source"] in {"APPROVED_MAPPING", "HISTORICAL_VOUCHER"}
    assert expense["reason"]


# --- Bank -> draft voucher -------------------------------------------------


def test_bank_receipt_and_payment_create_balanced_vouchers(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    _, _, receipt = _bank_ready(client, company["id"], credit="500.00", debit="0.00")
    receipt_response = client.post(
        f"{_base(company['id'])}/vouchers/from-bank/{receipt['id']}"
    )
    assert receipt_response.status_code == 200, receipt_response.text
    receipt_voucher = receipt_response.json()
    assert receipt_voucher["voucher_type"] == "RECEIPT"
    assert receipt_voucher["is_balanced"] is True
    assert any(
        line["ledger_name"] == "HDFC Current Account" and line["side"] == "Dr"
        for line in receipt_voucher["lines"]
    )
    assert any(
        line["ledger_name"] == "Deccan Freight Carriers" and line["side"] == "Cr"
        for line in receipt_voucher["lines"]
    )

    _, _, payment = _bank_ready(
        client,
        company["id"],
        credit="0.00",
        debit="250.00",
        reference_number="UTR-2",
    )
    payment_response = client.post(
        f"{_base(company['id'])}/vouchers/from-bank/{payment['id']}"
    )
    assert payment_response.status_code == 200, payment_response.text
    payment_voucher = payment_response.json()
    assert payment_voucher["voucher_type"] == "PAYMENT"
    assert payment_voucher["is_balanced"] is True
    assert any(
        line["ledger_name"] == "HDFC Current Account" and line["side"] == "Cr"
        for line in payment_voucher["lines"]
    )


def test_bank_source_voucher_creation_is_idempotent(client: TestClient) -> None:
    company, _ledgers = _company_with_ledgers(client)
    _, _, transaction = _bank_ready(client, company["id"], credit="500.00")
    first = client.post(
        f"{_base(company['id'])}/vouchers/from-bank/{transaction['id']}"
    )
    second = client.post(
        f"{_base(company['id'])}/vouchers/from-bank/{transaction['id']}"
    )
    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    assert first.json()["id"] == second.json()["id"]

    listed = client.get(
        f"{_base(company['id'])}/vouchers",
        params={"source_type": "BANK_TRANSACTION", "source_id": transaction["id"]},
    )
    assert listed.status_code == 200
    assert len(listed.json()) == 1


def test_bank_contra_flag_creates_contra_voucher(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    other_bank = _ledger(client, company["id"], "ICICI Current Account")
    _, _, transaction = _bank_ready(
        client,
        company["id"],
        credit="0.00",
        debit="300.00",
        counterparty_ledger="ICICI Current Account",
        is_inter_account_transfer=True,
    )
    response = client.post(
        f"{_base(company['id'])}/vouchers/from-bank/{transaction['id']}"
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert voucher["voucher_type"] == "CONTRA"
    assert voucher["is_balanced"] is True


def test_bank_narration_does_not_infer_an_inter_account_transfer(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    _ledger(client, company["id"], "ICICI Current Account")
    _, _, transaction = _bank_ready(
        client,
        company["id"],
        credit="0.00",
        debit="300.00",
        counterparty_ledger="ICICI Current Account",
        description="TRANSFER TO ICICI CURRENT ACCOUNT",
        is_inter_account_transfer=False,
    )
    response = client.post(
        f"{_base(company['id'])}/vouchers/from-bank/{transaction['id']}"
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert voucher["voucher_type"] == "PAYMENT"
    assert voucher["is_balanced"] is True


def test_bank_split_transaction_is_represented_by_multiple_counterparty_lines(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    _, _, transaction = _bank_ready(client, company["id"], credit="0.00", debit="300.00")
    response = client.post(
        f"{_base(company['id'])}/vouchers/from-bank/{transaction['id']}",
        json={
            "split_lines": [
                {"ledger_id": ledgers["expense"]["id"], "side": "Dr", "amount": "100.00"},
                {"ledger_id": ledgers["party"]["id"], "side": "Dr", "amount": "200.00"},
            ]
        },
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert len(voucher["lines"]) == 3
    assert voucher["is_balanced"] is True


def test_bank_mapping_update_is_company_scoped_and_audited(client: TestClient) -> None:
    company, _ledgers = _company_with_ledgers(client, name="Mapping Co")
    _account, statement, transaction = _bank_ready(
        client, company["id"], counterparty_ledger=None
    )
    other = _company(client, "Mapping Other Co")

    foreign = client.patch(
        f"{_base(other['id'])}/bank-statements/{statement['id']}"
        f"/transactions/{transaction['id']}/mapping",
        json={"counterparty_ledger": "Deccan Freight Carriers"},
    )
    assert foreign.status_code == 404

    updated = client.patch(
        f"{_base(company['id'])}/bank-statements/{statement['id']}"
        f"/transactions/{transaction['id']}/mapping",
        json={
            "counterparty_name": "Deccan Freight Carriers",
            "counterparty_ledger": "Deccan Freight Carriers",
            "is_inter_account_transfer": True,
        },
        headers={"X-PAL-Actor": "reviewer"},
    )
    assert updated.status_code == 200, updated.text
    body = updated.json()
    assert body["counterparty_ledger"] == "Deccan Freight Carriers"
    assert body["is_inter_account_transfer"] is True

    audit = client.get(
        f"{_base(company['id'])}/accounting-audit",
        params={"entity_type": "MAPPING"},
    )
    assert audit.status_code == 200
    events = audit.json()
    assert any(
        event["event_type"] == "BANK_TRANSACTION_MAPPING_UPDATED"
        and event["actor"] == "reviewer"
        and event["field"] == "counterparty_ledger"
        and event["new_value"] == "Deccan Freight Carriers"
        for event in events
    )
    assert any(
        event["event_type"] == "BANK_TRANSACTION_MAPPING_UPDATED"
        and event["field"] == "is_inter_account_transfer"
        and event["new_value"] == "true"
        for event in events
    )


def test_bank_missing_counterparty_routes_to_review(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    _, _, transaction = _bank_ready(
        client, company["id"], counterparty_ledger=None, credit="500.00"
    )
    response = client.post(
        f"{_base(company['id'])}/vouchers/from-bank/{transaction['id']}"
    )
    assert response.status_code == 200, response.text
    voucher = response.json()
    assert voucher["status"] == "NEEDS_REVIEW"
    assert any(
        issue["code"] in {"MISSING_COUNTERPARTY", "NO_HISTORICAL_MAPPING"}
        for issue in voucher["validation_issues"]
    )
    approved = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/approve", json={}
    )
    assert approved.status_code == 409
    assert approved.json()["detail"]["code"] == "INVALID_STATUS_TRANSITION"


# --- End-to-end source workflows ------------------------------------------


def test_invoice_and_bank_vouchers_complete_the_export_lifecycle(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    _ledger(client, company["id"], "ICICI Current Account")

    document = _document(client, company["id"])
    vendor = _vendor(client, company["id"])
    purchase_invoice = _invoice(
        client, company["id"], document["id"], vendor["id"], invoice_number="PUR-1"
    )
    customer = _vendor(
        client, company["id"], name="Acme Customer", expense_ledger=None
    )
    sales_invoice = _invoice(
        client,
        company["id"],
        document["id"],
        customer["id"],
        direction="SALES",
        invoice_number="SAL-1",
    )

    source_vouchers = [
        client.post(
            f"{_base(company['id'])}/vouchers/from-invoice/{purchase_invoice['id']}"
        ).json(),
        client.post(
            f"{_base(company['id'])}/vouchers/from-invoice/{sales_invoice['id']}"
        ).json(),
    ]
    for credit, debit, reference in (
        ("500.00", "0.00", "UTR-RECEIPT"),
        ("0.00", "250.00", "UTR-PAYMENT"),
    ):
        _, _, transaction = _bank_ready(
            client,
            company["id"],
            credit=credit,
            debit=debit,
            reference_number=reference,
        )
        source_vouchers.append(
            client.post(
                f"{_base(company['id'])}/vouchers/from-bank/{transaction['id']}"
            ).json()
        )
    _, _, contra_transaction = _bank_ready(
        client,
        company["id"],
        credit="0.00",
        debit="300.00",
        counterparty_ledger="ICICI Current Account",
        is_inter_account_transfer=True,
        reference_number="UTR-CONTRA",
    )
    source_vouchers.append(
        client.post(
            f"{_base(company['id'])}/vouchers/from-bank/{contra_transaction['id']}"
        ).json()
    )

    assert [voucher["voucher_type"] for voucher in source_vouchers] == [
        "PURCHASE",
        "SALES",
        "RECEIPT",
        "PAYMENT",
        "CONTRA",
    ]
    assert all(voucher["status"] == "DRAFT" for voucher in source_vouchers)
    assert all(voucher["is_balanced"] is True for voucher in source_vouchers)

    for voucher in source_vouchers:
        submitted = client.post(
            f"{_base(company['id'])}/vouchers/{voucher['id']}/submit"
        )
        assert submitted.status_code == 200, submitted.text
        assert submitted.json()["status"] == "PENDING_APPROVAL"

        approved = client.post(
            f"{_base(company['id'])}/vouchers/{voucher['id']}/approve", json={}
        )
        assert approved.status_code == 200, approved.text
        assert approved.json()["status"] == "APPROVED"

        exported = client.post(
            f"{_base(company['id'])}/vouchers/{voucher['id']}/export"
        )
        assert exported.status_code == 200, exported.text
        after = client.get(
            f"{_base(company['id'])}/vouchers/{voucher['id']}"
        ).json()
        assert after["status"] == "AWAITING_IMPORT_CONFIRMATION"
        assert after["posted_at"] is None
        assert after["posting_reference"] is None


# --- Export and posting verification --------------------------------------


def test_export_transitions_to_awaiting_not_posted_and_verification_sets_posted(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    client.post(f"{_base(company['id'])}/vouchers/{voucher['id']}/submit")
    client.post(f"{_base(company['id'])}/vouchers/{voucher['id']}/approve", json={})

    exported = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/export",
        json={"actor": "Asha"},
    )
    assert exported.status_code == 200, exported.text
    assert exported.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(io.BytesIO(exported.content))
    assert workbook.sheetnames == [
        "Accounting Voucher",
        "Accounting Voucher (Read Me)",
    ]
    sheet = workbook["Accounting Voucher"]
    assert tuple(cell.value for cell in sheet[1]) == ACCOUNTING_COLUMNS
    assert workbook["Accounting Voucher (Read Me)"].max_row > 400
    rows = [
        row
        for row in sheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in row)
    ]
    assert len(rows) == 2
    assert sum(row[6] for row in rows if row[7] == "Dr") == 100
    assert sum(row[6] for row in rows if row[7] == "Cr") == 100

    after = client.get(f"{_base(company['id'])}/vouchers/{voucher['id']}").json()
    assert after["status"] == "AWAITING_IMPORT_CONFIRMATION"
    assert after["posted_at"] is None
    assert after["export_count"] == 1
    assert after["exported_by"] == "Asha"

    no_reference = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/verify-posting",
        json={"verified": True},
    )
    assert no_reference.status_code == 409
    assert no_reference.json()["detail"]["code"] == "POSTING_NOT_VERIFIED"

    posted = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/verify-posting",
        json={
            "verified": True,
            "external_reference": "TALLY-VCH-001",
            "reason": "Imported and checked in Tally Prime.",
        },
    )
    assert posted.status_code == 200, posted.text
    body = posted.json()
    assert body["status"] == "POSTED"
    assert body["posting_reference"] == "TALLY-VCH-001"
    assert body["posted_at"] is not None

    audit = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    types = [event["event_type"] for event in audit]
    assert "VOUCHER_EXPORTED" in types
    assert "VOUCHER_POSTED_VERIFIED" in types


def test_posting_verification_failure_marks_posting_failed(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    client.post(f"{_base(company['id'])}/vouchers/{voucher['id']}/submit")
    client.post(f"{_base(company['id'])}/vouchers/{voucher['id']}/approve", json={})
    client.post(f"{_base(company['id'])}/vouchers/{voucher['id']}/export")
    failed = client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/verify-posting",
        json={"verified": False, "reason": "Tally rejected the voucher."},
    )
    assert failed.status_code == 200, failed.text
    assert failed.json()["status"] == "POSTING_FAILED"
    assert failed.json()["posting_error"] == "Tally rejected the voucher."

    audit = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    assert any(
        event["event_type"] == "VOUCHER_POSTING_FAILED"
        and event["previous_status"] == "AWAITING_IMPORT_CONFIRMATION"
        and event["new_status"] == "POSTING_FAILED"
        and event["reason"] == "Tally rejected the voucher."
        for event in audit
    )


def test_export_preview_only_counts_approved_vouchers(client: TestClient) -> None:
    company, ledgers = _company_with_ledgers(client)
    draft = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    preview = client.get(f"{_base(company['id'])}/vouchers/export/preview")
    assert preview.status_code == 200, preview.text
    assert preview.json()["rows"] == 0
    assert preview.json()["vouchers"] == 0
    assert draft["status"] == "DRAFT"


# --- Cross-company security ------------------------------------------------


def test_voucher_ledger_and_source_are_cross_company_isolated(client: TestClient) -> None:
    a, a_ledgers = _company_with_ledgers(client, name="A")
    b, b_ledgers = _company_with_ledgers(client, name="B")
    voucher = _manual_voucher(
        client, a["id"], a_ledgers["expense"]["id"], a_ledgers["party"]["id"]
    ).json()

    cross_get = client.get(f"{_base(b['id'])}/vouchers/{voucher['id']}")
    assert cross_get.status_code == 404

    cross_ledger = client.post(
        f"{_base(b['id'])}/vouchers",
        json={
            "voucher_date": "2026-04-01",
            "voucher_type": "JOURNAL",
            "lines": [
                {"ledger_id": a_ledgers["expense"]["id"], "side": "Dr", "amount": "10.00"},
                {"ledger_id": b_ledgers["party"]["id"], "side": "Cr", "amount": "10.00"},
            ],
        },
    )
    assert cross_ledger.status_code == 404

    document_a = _document(client, a["id"])
    vendor_a = _vendor(client, a["id"])
    invoice_a = _invoice(client, a["id"], document_a["id"], vendor_a["id"])
    cross_source = client.post(
        f"{_base(b['id'])}/vouchers/from-invoice/{invoice_a['id']}"
    )
    assert cross_source.status_code == 404

    _, _, transaction_a = _bank_ready(client, a["id"], credit="500.00")
    cross_bank_source = client.post(
        f"{_base(b['id'])}/vouchers/from-bank/{transaction_a['id']}"
    )
    assert cross_bank_source.status_code == 404


def test_audit_events_are_append_only_across_normal_operations(
    client: TestClient,
) -> None:
    company, ledgers = _company_with_ledgers(client)
    voucher = _manual_voucher(
        client, company["id"], ledgers["expense"]["id"], ledgers["party"]["id"]
    ).json()
    before = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    assert before

    assert client.patch(
        f"{_base(company['id'])}/vouchers/{voucher['id']}",
        json={"narration": "Reviewed narration"},
    ).status_code == 200
    assert client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/submit"
    ).status_code == 200
    assert client.post(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/approve", json={}
    ).status_code == 200

    after = client.get(
        f"{_base(company['id'])}/vouchers/{voucher['id']}/audit"
    ).json()
    after_ids = {event["id"] for event in after}
    assert {event["id"] for event in before} <= after_ids
    for event in before:
        assert event in after
    assert len(after) > len(before)


def test_accounting_mapping_is_company_scoped(client: TestClient) -> None:
    a, a_ledgers = _company_with_ledgers(client, name="A")
    b, b_ledgers = _company_with_ledgers(client, name="B")
    mapping = client.post(
        f"{_base(a['id'])}/accounting-mappings",
        json={
            "source_type": "BANK_TRANSACTION",
            "match_type": "COUNTERPARTY",
            "counterparty_name": "Deccan Freight Carriers",
            "voucher_type": "PAYMENT",
            "ledger_id": a_ledgers["expense"]["id"],
            "reason": "Approved by the partner.",
        },
    )
    assert mapping.status_code == 201, mapping.text
    mapping_id = mapping.json()["id"]

    assert client.get(f"{_base(b['id'])}/accounting-mappings/{mapping_id}").status_code == 404
    cross_ledger = client.post(
        f"{_base(b['id'])}/accounting-mappings",
        json={
            "source_type": "BANK_TRANSACTION",
            "match_type": "COUNTERPARTY",
            "counterparty_name": "Deccan Freight Carriers",
            "ledger_id": a_ledgers["expense"]["id"],
            "reason": "Cross-company ledger must be rejected.",
        },
    )
    assert cross_ledger.status_code == 409


def test_company_audit_endpoint_returns_only_own_events(client: TestClient) -> None:
    a, a_ledgers = _company_with_ledgers(client, name="A")
    b, _ = _company_with_ledgers(client, name="B")
    _manual_voucher(
        client, a["id"], a_ledgers["expense"]["id"], a_ledgers["party"]["id"]
    )
    patched = client.patch(
        f"{_base(a['id'])}",
        json={"sales_ledger": "Domestic Sales"},
        headers={"X-PAL-Actor": "Asha"},
    )
    assert patched.status_code == 200, patched.text

    a_audit = client.get(f"{_base(a['id'])}/accounting-audit")
    b_audit = client.get(f"{_base(b['id'])}/accounting-audit")
    assert a_audit.status_code == 200
    assert b_audit.status_code == 200
    assert a_audit.json()
    assert b_audit.json()
    assert all(event["company_id"] == b["id"] for event in b_audit.json())
    assert all(event["company_id"] == a["id"] for event in a_audit.json())
    assert any(
        event["event_type"] == "COMPANY_LEDGER_DEFAULT_UPDATED"
        and event["actor"] == "Asha"
        and event["field"] == "sales_ledger"
        and event["new_value"] == "Domestic Sales"
        for event in a_audit.json()
    )
