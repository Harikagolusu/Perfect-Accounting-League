"""Structured text extraction for XLSX and DOCX.

The readers are pure with respect to the application: they take bytes and return
bounded text. These tests verify both the values and the limits that keep a
large, unrelated document out of the prompt.
"""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from app.services.document_errors import DocumentPreparationError
from app.services.structured_documents import (
    MAX_CELL_CHARACTERS,
    MAX_DOCX_PARAGRAPHS,
    MAX_SPREADSHEET_COLUMNS,
    MAX_SPREADSHEET_ROWS,
    MAX_SPREADSHEET_SHEETS,
    MAX_STRUCTURED_CHARACTERS,
    TRUNCATION_MARKER,
    extract_structured_text,
)


def _xlsx(rows, *, title: str = "Invoice") -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _docx(paragraphs=(), tables=()) -> bytes:
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    for table in tables:
        created = document.add_table(rows=0, cols=max(len(row) for row in table))
        for row in table:
            cells = created.add_row().cells
            for index, value in enumerate(row):
                cells[index].text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_xlsx_renders_sheets_as_readable_text() -> None:
    raw = _xlsx(
        (
            ("TAX INVOICE",),
            ("Invoice No", "DFC/2026/0042"),
            ("Invoice Date", date(2026, 3, 15)),
            ("Taxable Amount", 1000),
            ("Total", 1180.5),
        )
    )

    result = extract_structured_text(".xlsx", raw)

    assert result.truncated is False
    assert "--- Sheet: Invoice ---" in result.text
    assert "Invoice No | DFC/2026/0042" in result.text
    assert "2026-03-15" in result.text
    assert "1000" in result.text
    assert "1180.5" in result.text


def test_xlsx_skips_blank_rows_and_trailing_empty_cells() -> None:
    raw = _xlsx(
        (
            ("Item", "", ""),
            ("", "", ""),
            ("Widget", "", ""),
        )
    )

    result = extract_structured_text(".xlsx", raw)

    assert result.text.count("Widget") == 1
    assert "Item |" not in result.text
    assert result.text.count("\n") >= 2


def test_xlsx_row_limit_truncates_and_marks_the_result() -> None:
    rows = (("Header", "Value"),) + tuple(
        (f"Row {index}", str(index)) for index in range(MAX_SPREADSHEET_ROWS + 20)
    )

    result = extract_structured_text(".xlsx", _xlsx(rows))

    assert result.truncated is True
    assert TRUNCATION_MARKER.strip() in result.text
    assert f"Row {MAX_SPREADSHEET_ROWS + 10}" not in result.text


def test_xlsx_multiple_sheets_keep_order_and_enforce_the_sheet_limit() -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index in range(MAX_SPREADSHEET_SHEETS + 2):
        sheet = workbook.create_sheet(title=f"Sheet{index + 1}")
        sheet.append(("Invoice No", f"INV-{index + 1}"))
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    result = extract_structured_text(".xlsx", buffer.getvalue())

    assert result.truncated is True
    for index in range(MAX_SPREADSHEET_SHEETS):
        assert f"--- Sheet: Sheet{index + 1} ---" in result.text
    assert "Sheet6" not in result.text
    assert result.text.index("INV-1") < result.text.index("INV-2")


def test_xlsx_column_limit_truncates_and_marks_the_result() -> None:
    row = tuple(f"col-{index}" for index in range(MAX_SPREADSHEET_COLUMNS + 5))
    result = extract_structured_text(".xlsx", _xlsx((row,)))

    assert result.truncated is True
    assert TRUNCATION_MARKER.strip() in result.text
    assert f"col-{MAX_SPREADSHEET_COLUMNS}" not in result.text


def test_xlsx_one_enormous_cell_is_clipped_before_assembly() -> None:
    huge = "x" * (MAX_CELL_CHARACTERS + 100)
    result = extract_structured_text(".xlsx", _xlsx(((huge,),)))

    assert result.truncated is True
    assert TRUNCATION_MARKER.strip() in result.text
    assert len(result.text) < MAX_CELL_CHARACTERS + len(TRUNCATION_MARKER) + 100


def test_xlsx_later_sheet_survives_a_row_limit_on_the_first_sheet() -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "Ledger"
    for index in range(MAX_SPREADSHEET_ROWS + 5):
        first.append((f"line-{index}",))
    second = workbook.create_sheet(title="Invoice")
    second.append(("Invoice No", "DFC/2026/0042"))
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    result = extract_structured_text(".xlsx", buffer.getvalue())

    assert result.truncated is True
    assert "--- Sheet: Invoice ---" in result.text
    assert "Invoice No | DFC/2026/0042" in result.text


def test_xlsx_formulas_are_not_evaluated_or_echoed() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Total"
    sheet["B1"] = "=1+1"
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()

    result = extract_structured_text(".xlsx", buffer.getvalue())

    # openpyxl in data_only mode has no cached formula result to show. The
    # formula expression itself must never reach the prompt.
    assert "=1+1" not in result.text
    assert result.text == "--- Sheet: Sheet ---\nTotal"


def test_xlsx_character_limit_is_enforced() -> None:
    rows = (("Header", "Value"),) + tuple(
        ("x" * 500, "y" * 500) for _ in range(200)
    )

    result = extract_structured_text(".xlsx", _xlsx(rows))

    assert result.truncated is True
    assert len(result.text) <= MAX_STRUCTURED_CHARACTERS + len(TRUNCATION_MARKER) + 1


def test_docx_reads_paragraphs_and_tables() -> None:
    raw = _docx(
        paragraphs=(
            "TAX INVOICE",
            "Deccan Freight Carriers",
            "Invoice No: DFC/2026/0042",
        ),
        tables=(
            (
                ("Taxable Amount", "1000.00"),
                ("Grand Total", "1180.00"),
            ),
        ),
    )

    result = extract_structured_text(".docx", raw)

    assert result.truncated is False
    assert "Deccan Freight Carriers" in result.text
    assert "Invoice No: DFC/2026/0042" in result.text
    assert "--- Table ---" in result.text
    assert "Taxable Amount | 1000.00" in result.text


def test_docx_table_only_invoice_is_not_ignored() -> None:
    raw = _docx(
        tables=(
            (
                ("TAX INVOICE",),
                ("Invoice No", "DFC/2026/0042"),
                ("Grand Total", "1180.00"),
            ),
        )
    )

    result = extract_structured_text(".docx", raw)

    assert "--- Table ---" in result.text
    assert "Invoice No | DFC/2026/0042" in result.text
    assert "Grand Total | 1180.00" in result.text


def test_docx_one_enormous_paragraph_is_clipped_before_assembly() -> None:
    huge = "y" * (MAX_CELL_CHARACTERS + 100)
    result = extract_structured_text(".docx", _docx(paragraphs=(huge,)))

    assert result.truncated is True
    assert TRUNCATION_MARKER.strip() in result.text
    assert len(result.text) < MAX_CELL_CHARACTERS + len(TRUNCATION_MARKER) + 100


def test_docx_paragraph_limit_is_explicit() -> None:
    raw = _docx(
        paragraphs=tuple(
            f"Invoice paragraph {index}"
            for index in range(MAX_DOCX_PARAGRAPHS + 5)
        )
    )

    result = extract_structured_text(".docx", raw)

    assert result.truncated is True
    assert TRUNCATION_MARKER.strip() in result.text
    assert f"Invoice paragraph {MAX_DOCX_PARAGRAPHS + 4}" not in result.text


def test_empty_documents_still_return_empty_text() -> None:
    # An empty workbook always has one empty sheet; it must not look like
    # readable invoice content.
    result = extract_structured_text(".xlsx", _xlsx((("",),)))

    assert result.text == ""
    assert result.truncated is False


def test_empty_docx_has_no_text() -> None:
    result = extract_structured_text(".docx", _docx())

    assert result.text == ""
    assert result.truncated is False


def test_legacy_xls_is_rejected_with_a_safe_message() -> None:
    with pytest.raises(DocumentPreparationError) as caught:
        extract_structured_text(".xls", b"\xd0\xcf\x11\xe0")

    assert caught.value.code == DocumentPreparationError.DOCUMENT_UNREADABLE
    assert ".xlsx" in caught.value.message


def test_corrupt_xlsx_is_unreadable() -> None:
    with pytest.raises(DocumentPreparationError) as caught:
        extract_structured_text(".xlsx", b"PK\x03\x04not a real workbook")

    assert caught.value.code == DocumentPreparationError.DOCUMENT_UNREADABLE
    assert "spreadsheet" in caught.value.message.lower()


def test_corrupt_docx_is_unreadable() -> None:
    with pytest.raises(DocumentPreparationError) as caught:
        extract_structured_text(".docx", b"PK\x03\x04not a real document")

    assert caught.value.code == DocumentPreparationError.DOCUMENT_UNREADABLE
    assert "docx" in caught.value.message.lower()


def test_archive_expansion_limit_is_checked_before_parsing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "app.services.structured_documents.MAX_UNCOMPRESSED_STRUCTURED_BYTES", 10
    )

    with pytest.raises(DocumentPreparationError) as caught:
        extract_structured_text(".xlsx", _xlsx((("Header",),)))

    assert caught.value.code == DocumentPreparationError.DOCUMENT_TOO_LARGE


def test_archive_entry_limit_is_checked_before_parsing(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        "app.services.structured_documents.MAX_ARCHIVE_ENTRIES", 1
    )

    with pytest.raises(DocumentPreparationError) as caught:
        extract_structured_text(".xlsx", _xlsx((("Header",),)))

    assert caught.value.code == DocumentPreparationError.DOCUMENT_TOO_LARGE


def test_unexpected_extension_is_rejected() -> None:
    with pytest.raises(DocumentPreparationError) as caught:
        extract_structured_text(".csv", b"a,b,c")

    assert caught.value.code == DocumentPreparationError.UNSUPPORTED_DOCUMENT_TYPE


def test_datetime_cells_are_rendered_unambiguously() -> None:
    raw = _xlsx((("Timestamp", datetime(2026, 3, 15, 13, 45)),))

    result = extract_structured_text(".xlsx", raw)

    assert "2026-03-15 13:45:00" in result.text
